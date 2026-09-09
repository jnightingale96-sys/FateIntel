from __future__ import annotations

import os
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "FATEINTEL_NANO_PFAS_REGULATORY_REQUIREMENTS.md"
OUTPUT = ROOT.parent / "FateIntel_Nano_PFAS_Regulatory_Fate_Assessment_Report.docx"

NAVY = "173A46"
TEAL = "00877C"
PALE_TEAL = "EAF6F4"
PALE_BLUE = "EAF0F2"
MID_GREY = "68777D"
LIGHT_GREY = "F3F5F5"
WHITE = "FFFFFF"
BLACK = "20282B"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        edge_data = edges[edge]
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_repeat_title_rows(table) -> None:
    if table.rows:
        set_repeat_table_header(table.rows[0])


def set_table_widths(table) -> None:
    """Set explicit fixed widths because LibreOffice ignores Pandoc's bare tblGrid."""

    width_map = {
        2: (0.31, 0.69),
        3: (0.18, 0.31, 0.51),
        4: (0.14, 0.24, 0.38, 0.24),
    }
    ratios = width_map.get(len(table.columns), tuple(1 / len(table.columns) for _ in table.columns))
    total_twips = 10000
    widths = [round(total_twips * ratio) for ratio in ratios]
    widths[-1] += total_twips - sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_twips))
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_cols = list(table._tbl.tblGrid.gridCol_lst)
    for index, width in enumerate(widths):
        if index < len(grid_cols):
            grid_cols[index].set(qn("w:w"), str(width))
        table.columns[index].width = Twips(width)
        for cell in table.columns[index].cells:
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def rebuild_pandoc_tables(doc: Document) -> None:
    """Replace Pandoc tables with native python-docx tables for portable rendering.

    Pandoc's minimal table XML opens in Word but LibreOffice lays the cell text
    outside the grid.  Rebuilding retains the semantic rows and produces a DOCX
    that renders consistently in both applications.
    """

    for old in list(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in old.rows]
        if not rows:
            continue
        replacement = doc.add_table(rows=len(rows), cols=len(rows[0]))
        replacement.style = "Table Grid"
        for row_index, values in enumerate(rows):
            for column_index, value in enumerate(values):
                replacement.cell(row_index, column_index).text = value
        old._tbl.addprevious(replacement._tbl)
        old._tbl.getparent().remove(old._tbl)


def add_page_number(paragraph, *, align_right: bool = True) -> None:
    if align_right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def add_rule(paragraph, color: str = TEAL, size: int = 12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def font_for(style, name: str = "Aptos") -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def make_reference_doc(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.74)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = False

    styles = doc.styles
    normal = styles["Normal"]
    font_for(normal)
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    font_for(title)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(98)
    title.paragraph_format.space_after = Pt(13)
    title.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    font_for(subtitle)
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(TEAL)
    subtitle.font.italic = False
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.keep_with_next = True

    if "Author" in styles:
        author = styles["Author"]
        font_for(author)
        author.font.size = Pt(10.5)
        author.font.bold = True
        author.font.color.rgb = RGBColor.from_string(MID_GREY)

    if "Date" in styles:
        date = styles["Date"]
        font_for(date)
        date.font.size = Pt(10)
        date.font.color.rgb = RGBColor.from_string(MID_GREY)
        date.paragraph_format.space_before = Pt(6)

    heading_specs = {
        "Heading 1": (18, NAVY, 16, 6),
        "Heading 2": (13.5, TEAL, 12, 4),
        "Heading 3": (11.2, NAVY, 9, 3),
        "Heading 4": (10.2, MID_GREY, 7, 2),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        font_for(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    styles["Heading 1"].paragraph_format.page_break_before = True

    for name in ("TOC Heading", "Contents Heading"):
        if name in styles:
            style = styles[name]
            font_for(style)
            style.font.size = Pt(20)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(NAVY)
            style.paragraph_format.page_break_before = True
            style.paragraph_format.space_after = Pt(14)

    for name in ("TOC 1", "TOC 2", "TOC 3"):
        if name in styles:
            style = styles[name]
            font_for(style)
            style.font.size = Pt(9.2 if name != "TOC 1" else 9.7)
            style.font.color.rgb = RGBColor.from_string(BLACK)
            style.paragraph_format.space_after = Pt(2)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        font_for(style)
        style.font.size = Pt(9.5)
        style.paragraph_format.space_after = Pt(2.5)
        style.paragraph_format.left_indent = Inches(0.22)
        style.paragraph_format.first_line_indent = Inches(-0.16)

    quote = styles["Quote"]
    font_for(quote)
    quote.font.size = Pt(9.3)
    quote.font.italic = False
    quote.font.color.rgb = RGBColor.from_string(NAVY)
    quote.paragraph_format.left_indent = Inches(0.23)
    quote.paragraph_format.right_indent = Inches(0.18)
    quote.paragraph_format.space_before = Pt(5)
    quote.paragraph_format.space_after = Pt(7)

    if "Caption" in styles:
        caption = styles["Caption"]
        font_for(caption)
        caption.font.size = Pt(8.5)
        caption.font.bold = True
        caption.font.color.rgb = RGBColor.from_string(MID_GREY)

    if "Hyperlink" in styles:
        hyperlink = styles["Hyperlink"]
        font_for(hyperlink)
        hyperlink.font.color.rgb = RGBColor.from_string(TEAL)
        hyperlink.font.underline = True

    if "Source Code" in styles:
        code = styles["Source Code"]
        font_for(code, "Aptos Mono")
        code.font.size = Pt(8.3)
        code.font.color.rgb = RGBColor.from_string(NAVY)
        code.paragraph_format.left_indent = Inches(0.18)
        code.paragraph_format.space_after = Pt(5)

    if "Source note" not in styles:
        source_note = styles.add_style("Source note", WD_STYLE_TYPE.PARAGRAPH)
        font_for(source_note)
        source_note.font.size = Pt(8.4)
        source_note.font.color.rgb = RGBColor.from_string(MID_GREY)
        source_note.paragraph_format.space_after = Pt(3)

    core = doc.core_properties
    core.title = "Regulatory Environmental Fate and Ecotoxicology Assessment Requirements for Nanoforms, Nanomaterials and PFAS"
    core.subject = "FateIntel software specification and jurisdictional comparison"
    core.author = "FateIntel"
    core.comments = "Private working report; regulatory source status must be rechecked before use."
    doc.save(path)


def source_for_pandoc() -> str:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    while lines and (not lines[0].strip() or lines[0].startswith("**FateIntel") or lines[0].startswith("**Evidence") or lines[0].startswith("**Implementation")):
        lines = lines[1:]
    adjusted: list[str] = []
    for line in lines:
        if line.startswith("#### "):
            line = "### " + line[5:]
        elif line.startswith("### "):
            line = "## " + line[4:]
        elif line.startswith("## "):
            line = "# " + line[3:]
        bullet = re.match(r"^(\s*)-\s+(.*)$", line)
        numbered = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if bullet:
            line = f"{bullet.group(1)}• {bullet.group(2)}"
        elif numbered:
            line = f"{numbered.group(1)}{numbered.group(2)}) {numbered.group(3)}"
        adjusted.append(line)
        if bullet or numbered:
            adjusted.append("")
    text = "\n".join(adjusted)
    workflow = """**Workflow sequence**

| Step | Decision or record |
| --- | --- |
| 1 | Establish identity, jurisdiction and specialist-group applicability. |
| 2 | Build the nano material-state graph or PFAS member/precursor/product graph. |
| 3 | Reconcile exposure and effect concentration bases. |
| 4 | If a verified regulatory method exists, create an evidence-bound tier decision. |
| 5 | Otherwise require specialist review or a controlled external-model handoff. |
"""
    text = re.sub(r"```mermaid\n.*?\n```", workflow, text, flags=re.DOTALL)
    front_matter = """---
title: "Regulatory Environmental Fate and Ecotoxicology Assessment Requirements for Nanoforms, Nanomaterials and PFAS"
subtitle: "FateIntel software specification and jurisdictional comparison"
author: "Private working report · regulatory evidence review"
date: "Evidence cut-off: 9 September 2026"
---

"""
    contents = """# Contents

• Executive conclusion

• 1. Scope and status language

• 2. Nanoforms and nanomaterials

• 3. PFAS and associated chemicals

• 4. Cross-model compatibility rules

• 5. Proposed FateIntel workflow

• 6. What should be built now—and what should wait

• 7. Limitations and change control

• Sources

The contents are intentionally concise. Detailed subsections appear under each major heading.

"""
    return front_matter + contents + text + "\n"


def style_output(path: Path) -> None:
    doc = Document(path)
    rebuild_pandoc_tables(doc)
    # Populate both odd/default and even header/footer parts explicitly.  Some
    # Word-compatible renderers otherwise create an unstyled even-page part
    # during pagination, which makes alternate pages lose the running header.
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.74)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.27)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = False

    def populate_header(header) -> None:
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = hp.add_run("FATEINTEL  /  REGULATORY EVIDENCE & SOFTWARE SPECIFICATION")
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
        add_rule(hp, PALE_BLUE, 6)

    def populate_footer(footer) -> None:
        footer.is_linked_to_previous = False
        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.clear()
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        footer_paragraph.paragraph_format.space_after = Pt(0)
        tabs = OxmlElement("w:tabs")
        right_tab = OxmlElement("w:tab")
        right_tab.set(qn("w:val"), "right")
        right_tab.set(qn("w:pos"), "10080")  # 7.0 inches, the usable text width.
        tabs.append(right_tab)
        footer_paragraph._p.get_or_add_pPr().append(tabs)
        left = footer_paragraph.add_run(
            "PRIVATE WORKING REPORT  ·  SOURCE STATUS MUST BE RECHECKED BEFORE REGULATORY USE"
        )
        left.font.name = "Aptos"
        left.font.size = Pt(6.8)
        left.font.color.rgb = RGBColor.from_string(MID_GREY)
        footer_paragraph.add_run("\t")
        add_page_number(footer_paragraph, align_right=False)

    populate_header(section.header)
    populate_footer(section.footer)
    populate_header(section.even_page_header)
    populate_footer(section.even_page_footer)

    title_seen = False
    in_sources = False
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style else ""
        if name == "Title":
            title_seen = True
            add_rule(paragraph, TEAL, 18)
        elif title_seen and name == "Subtitle":
            paragraph.paragraph_format.space_before = Pt(6)
        if paragraph.text.strip() == "3.5 PFAS routes by jurisdiction":
            paragraph.paragraph_format.page_break_before = True
        if paragraph.text.strip() == "Sources" and name.startswith("Heading"):
            in_sources = True
        elif in_sources and paragraph.text.strip() and not name.startswith("Heading"):
            paragraph.style = doc.styles["Source note"]
        if paragraph.text.startswith("This specification does **not**"):
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.right_indent = Inches(0.22)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(8)
            add_rule(paragraph, TEAL, 8)
        stripped = paragraph.text.lstrip()
        if stripped.startswith("• ") or re.match(r"^\d+\)\s", stripped):
            paragraph.paragraph_format.left_indent = Inches(0.23)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            paragraph.paragraph_format.space_after = Pt(2.5)

    border = {"val": "single", "sz": "4", "space": "0", "color": "D9E2E3"}
    for table in doc.tables:
        table.style = "Table Grid"
        table.autofit = True
        set_table_widths(table)
        set_repeat_title_rows(table)
        for row_idx, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_margins(cell)
                set_cell_border(cell, top=border, left=border, bottom=border, right=border)
                shade(cell, NAVY if row_idx == 0 else (LIGHT_GREY if row_idx % 2 == 0 else WHITE))
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1.5)
                    paragraph.paragraph_format.keep_together = True
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        run.font.size = Pt(8.2)
                        run.font.color.rgb = RGBColor.from_string(WHITE if row_idx == 0 else BLACK)
                        if row_idx == 0:
                            run.font.bold = True

    doc.core_properties.title = "Regulatory Environmental Fate and Ecotoxicology Assessment Requirements for Nanoforms, Nanomaterials and PFAS"
    doc.core_properties.subject = "FateIntel software specification and jurisdictional comparison"
    doc.core_properties.author = "FateIntel"
    doc.save(path)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="fateintel-docx-") as temp_dir:
        temp = Path(temp_dir)
        reference = temp / "reference.docx"
        source = temp / "report.md"
        draft = temp / "draft.docx"
        make_reference_doc(reference)
        source.write_text(source_for_pandoc(), encoding="utf-8")
        env = dict(os.environ)
        subprocess.run(
            [
                "pandoc",
                str(source),
                "--from=gfm",
                "--to=docx",
                "--reference-doc",
                str(reference),
                "--output",
                str(draft),
            ],
            check=True,
            env=env,
        )
        draft.replace(OUTPUT)
    style_output(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
